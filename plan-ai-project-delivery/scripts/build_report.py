#!/usr/bin/env python3
"""Build a review-ready *layout draft* from a deliberately small Markdown subset.

Dependencies: python-docx, Pillow. Rendering and all-page visual QA remain required.
See ../references/word-delivery.md for input contract and examples.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import os
from urllib.parse import unquote, quote, urlparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


class InputError(ValueError):
    pass


IMAGE = re.compile(r"!\[([^\]]+)\]\(([^)]+)\)$")
FIGURE = re.compile(r"<!-- figure-page: (a3|a4-landscape) -->$")
INLINE = re.compile(r"(\*\*[^*\n]+\*\*|`[^`\n]+`|\[[^\]\n]+\]\([^)\s]+\))")


def inline_parts(text):
    """Reject unhandled formatting rather than silently deleting it."""
    parts = INLINE.split(text)
    for idx, value in enumerate(parts):
        if idx % 2 == 0 and re.search(r"\*|`|!\[|\[[^\]]+\]\(|\[\^|~~|<[^>]+>|(?<!\w)_{1,2}\S.*?_{1,2}(?!\w)", value):
            raise InputError(f"Unsupported inline Markdown: {text[:100]}")
    return parts


def cells(line):
    if not line.endswith('|'):
        raise InputError('Table rows must start and end with |')
    return [x.strip().replace(r'\|', '|') for x in re.split(r'(?<!\\)\|', line[1:-1])]


def parse(source: Path):
    lines = source.read_text(encoding='utf-8').splitlines()
    result = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        number = i + 1
        if raw.startswith((' ', '\t')):
            raise InputError(f'Line {number}: indentation/nested lists are unsupported')
        figure = FIGURE.fullmatch(line)
        if figure:
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i >= len(lines) or not IMAGE.fullmatch(lines[i].strip()):
                raise InputError(f'Line {number}: figure-page must precede one image')
            line = lines[i].strip()
        image = IMAGE.fullmatch(line)
        if image:
            caption, relative = image.groups()
            if relative.startswith(('http:', 'https:', 'data:', 'file:')) or Path(relative).is_absolute():
                raise InputError(f'Line {number}: image must use a local relative path')
            path = (source.parent / relative).resolve()
            if not path.is_file():
                raise InputError(f'Line {number}: missing image {relative}')
            with Image.open(path) as im:
                width, height = im.size
                if im.format not in ('PNG', 'JPEG'):
                    raise InputError(f'Line {number}: use PNG or JPEG, not {im.format}')
                im.verify()
            result.append(('image', caption, path, width, height, figure.group(1) if figure else None))
        elif re.match(r'^#{1,3} ', line):
            marks, title = line.split(' ', 1)
            if not title.strip():
                raise InputError(f'Line {number}: empty heading')
            inline_parts(title)
            result.append(('heading', len(marks), title))
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(cells(lines[i].strip()))
                i += 1
            if len(rows) < 2 or not all(re.fullmatch(r':?-{3,}:?', x) for x in rows[1]):
                raise InputError(f'Line {number}: missing Markdown table separator')
            width = len(rows[0])
            if width > 6 or any(len(row) != width for row in rows):
                raise InputError(f'Line {number}: table widths disagree or exceed 6 columns')
            rows.pop(1)
            for row in rows:
                for value in row:
                    inline_parts(value)
            result.append(('table', rows))
            continue
        elif re.match(r'^(- |\d+\. )', line):
            prefix, value = line.split(' ', 1)
            if re.match(r'^\[[ xX]\]', value):
                raise InputError(f'Line {number}: task checkboxes unsupported')
            inline_parts(value)
            result.append(('list', prefix, value))
        elif line == '<!-- page-break -->':
            result.append(('break',))
        elif re.match(r'^(#|>|```|~~~|<!--|[*+] |\*{3,}$|---+$|___+$)', line):
            raise InputError(f'Line {number}: unsupported block {line[:50]}')
        else:
            paragraph = [line]
            i += 1
            while i < len(lines) and lines[i].strip():
                following = lines[i]
                if re.match(r'^(\s|#|\||!\[|<!--|- |\d+\. |>|```|~~~|[*+] |\*{3,}$|---+$|___+$)', following):
                    break
                paragraph.append(following.strip())
                i += 1
            text = ' '.join(paragraph)
            inline_parts(text)
            result.append(('paragraph', text))
            continue
        i += 1
    headings = [b for b in result if b[0] == 'heading']
    if not headings or headings[0][1] != 1:
        raise InputError('Body must begin its heading hierarchy with # chapter')
    previous = 0
    for _, level, _ in headings:
        if level > previous + 1:
            raise InputError('Heading hierarchy must not skip levels')
        previous = level
    return result


def element(tag, **attrs):
    item = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        item.set(qn('w:' + key), str(value))
    return item


def font(style, family, size):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), family)
    for key in list(fonts.attrib):
        if key.rsplit('}', 1)[-1].endswith('Theme'):
            del fonts.attrib[key]


def geometry(section, kind=None):
    width, height = {'a3': (42, 29.7), 'a4-landscape': (29.7, 21)}.get(kind, (21, 29.7))
    section.orientation = WD_ORIENT.LANDSCAPE if kind else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(width), Cm(height)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2 if kind else 2.5)
    section.header_distance = section.footer_distance = Cm(1)


def rich(paragraph, text, bold=None):
    for value in inline_parts(text):
        if not value:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', value)
        if link:
            label, url = link.groups()
            hyperlink = element('hyperlink')
            if url.startswith('#'):
                hyperlink.set(qn('w:anchor'), unquote(url[1:]))
            else:
                parsed = urlparse(url)
                if parsed.scheme not in ('http', 'https', 'mailto'):
                    if parsed.scheme or Path(unquote(parsed.path)).is_absolute():
                        raise InputError('Only portable relative attachment links are allowed')
                    source_root = paragraph.part._source_root
                    target = (source_root / unquote(parsed.path)).resolve()
                    if not target.is_file() or not target.is_relative_to(source_root):
                        raise InputError('Missing or out-of-client-directory attachment: '+url)
                    url = quote(os.path.relpath(target, paragraph.part._output_root).replace(os.sep,'/'))
                    paragraph.part._attachment_links.append({'target':url,'sha256':hashlib.sha256(target.read_bytes()).hexdigest()})
                hyperlink.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = element('r')
            props = element('rPr')
            props.append(element('u', val='single'))
            run.append(props)
            txt = element('t')
            txt.text = label
            run.append(txt)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            strong = value.startswith('**')
            code = value.startswith('`')
            run = paragraph.add_run(value[2:-2] if strong else value[1:-1] if code else value)
            run.bold = True if strong else bold


def table(doc, rows):
    tab = doc.add_table(rows=0, cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    weights=[]
    for index, header in enumerate(rows[0]):
        values=[row[index] for row in rows[1:]]
        mean=sum(min(len(value),60) for value in values)/max(1,len(values))
        narrow=bool(re.search('ID|Share|Amount|Unit price|Tax rate|Start|End|Days',header)) and mean < 18
        weights.append(max(1.3,min(2.5,len(header)*.3+.7)) if narrow else max(2,min(7,mean**.5*.8)))
    section=doc.sections[-1]
    width_cm=(section.page_width-section.left_margin-section.right_margin)/360000
    widths=[width_cm*w/sum(weights) for w in weights]
    if rows[0] == ['Item', 'Content']:
        widths = [width_cm * .25, width_cm * .75]
    for col,width in zip(tab.columns,widths): col.width=Cm(width)
    for idx, values in enumerate(rows):
        row = tab.add_row()
        row._tr.get_or_add_trPr().append(element('cantSplit'))
        if idx == 0:
            row._tr.get_or_add_trPr().append(element('tblHeader', val='true'))
        for index,(cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            if idx == 0:
                props.append(element('shd', fill='EEEEEE'))
            margins = element('tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                margins.append(element(side, w='90', type='dxa'))
            props.append(margins)
            p = cell.paragraphs[0]
            p.style = 'Table Text'
            p.paragraph_format.keep_with_next = idx == 0
            if idx == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif re.fullmatch(r'[\d,.%％/—–\- ]+',value): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rich(p, value, bold=idx == 0)
    borders = element('tblBorders')
    for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        borders.append(element(side, val='single', sz=4, color='D9D9D9'))
    tab._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def field(paragraph, instruction, size=9):
    field_node = element('fldSimple', instr=' ' + instruction + ' ')
    run = element('r'); props = element('rPr')
    props.append(element('sz', val=int(size * 2)))
    props.append(element('szCs', val=int(size * 2)))
    run.append(props); field_node.append(run)
    paragraph._p.append(field_node)  # Layout engine calculates PAGE.


def page_numbering(section, start=None, fmt='decimal'):
    for node in list(section._sectPr.findall(qn('w:pgNumType'))):
        section._sectPr.remove(node)
    attrs = {'fmt': fmt}
    if start is not None: attrs['start'] = start
    section._sectPr.append(element('pgNumType', **attrs))


def footer_for(section, numbered=True):
    section.different_first_page_header_footer = False
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.style = 'Footer'
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if numbered:
        footer.add_run('Page ')
        field(footer, 'PAGE')
        footer.add_run(' ')
        for run in footer.runs: run.font.size = Pt(9)


def native_toc(doc, headings):
    """A genuine Word TOC field; finalizer rebuilds its cached result from headings."""
    sdt = element('sdt'); props = element('sdtPr')
    props.append(element('tag', val='report-native-toc'))
    gallery = element('docPartObj'); gallery.append(element('docPartGallery', val='Table of Contents'))
    props.append(gallery); sdt.append(props); content = element('sdtContent'); sdt.append(content)
    p = element('p'); ppr = element('pPr'); ppr.append(element('pStyle', val='TOC1')); p.append(ppr)
    for kind, text in [('begin', None), ('code', ' TOC \\o "1-2" \\h \\z '), ('separate', None), ('text', 'Contents pending generation'), ('end', None)]:
        r = element('r')
        node = element('instrText') if kind == 'code' else element('t') if kind == 'text' else element('fldChar', fldCharType=kind)
        if text: node.text = text
        r.append(node); p.append(r)
    content.append(p)
    doc._body._body.insert(len(doc._body._body)-1, sdt)


def build(args):
    source = args.markdown.resolve()
    if args.output.suffix.lower() != '.docx':
        raise InputError('Output must have .docx extension')
    metadata = json.loads(args.metadata.read_text(encoding='utf-8'))
    allowed = {'title', 'subtitle', 'version', 'status', 'date', 'audience', 'purpose', 'document_control', 'layout'}
    if not isinstance(metadata, dict) or set(metadata) - allowed:
        raise InputError('Metadata must be an object containing only documented keys')
    for key in ('title', 'version', 'status', 'date', 'purpose'):
        if not isinstance(metadata.get(key), str) or not metadata[key].strip():
            raise InputError(f'Metadata needs nonempty string: {key}')
    for key in ('subtitle', 'audience'):
        if key in metadata and not isinstance(metadata[key], str):
            raise InputError(f'Metadata {key} must be a string')
    controls = metadata.get('document_control', [])
    if not isinstance(controls, list) or any(not isinstance(r, list) or len(r) != 2 or any(not isinstance(v, str) for v in r) for r in controls):
        raise InputError('document_control must be a list of two-string rows')
    layout = metadata.get('layout', {})
    if not isinstance(layout, dict) or set(layout) - {'chapter_page_break'}:
        raise InputError('layout supports only chapter_page_break')
    chapter_break = layout.get('chapter_page_break', True)
    if not isinstance(chapter_break, bool):
        raise InputError('layout.chapter_page_break must be boolean')
    blocks = parse(source)
    doc = Document()
    doc.part._source_root = source.parent
    doc.part._output_root = args.output.resolve().parent
    doc.part._attachment_links = []
    # Some runtime default templates add decorative borders to Title styles.
    # This builder intentionally uses plain, black formal-document headings.
    for style in doc.styles:
        for border in list(style._element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    geometry(doc.sections[0])
    for name, family, size in [('Normal', args.body_font, 12), ('Title', args.heading_font, 24), ('Subtitle', args.heading_font, 16), ('Heading 1', args.heading_font, 17), ('Heading 2', args.heading_font, 14), ('Heading 3', args.heading_font, 12), ('Caption', args.body_font, 10.5), ('Footer', args.body_font, 9)]:
        st = doc.styles[name]
        font(st, family, size)
        st.font.italic = False
        st.paragraph_format.line_spacing = Pt(size * 1.5)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.widow_control = True
        if name.startswith('Heading'):
            st.font.bold = True
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(12)
    table_style = doc.styles.add_style('Table Text', 1)
    table_style.base_style = doc.styles['Normal']
    font(table_style, args.body_font, 10.5)
    table_style.paragraph_format.line_spacing = Pt(15)
    table_style.paragraph_format.space_after = Pt(0)
    for level in (1, 2):
        name = f'TOC {level}'
        st = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
        st.base_style = doc.styles['Normal']
        font(st, args.body_font, 12 if level == 1 else 11)
        st.font.bold = level == 1
        st.paragraph_format.line_spacing = Pt(19)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.left_indent = Cm((level - 1) * .6)
        st.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    front_heading = doc.styles.add_style('Front Heading', 1)
    front_heading.base_style = doc.styles['Normal']
    font(front_heading, args.heading_font, 14)
    front_heading.font.bold = True
    front_heading.paragraph_format.keep_with_next = True
    footer_for(doc.sections[0], numbered=False)
    doc.settings.element.append(element('updateFields', val='true'))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(66)
    for key, style in [('title', 'Title'), ('subtitle', 'Subtitle')]:
        if metadata.get(key):
            p = doc.add_paragraph(metadata[key], style)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for key in ('version', 'status', 'date', 'audience'):
        if metadata.get(key):
            p = doc.add_paragraph(metadata[key])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front = doc.add_section(WD_SECTION.NEW_PAGE)
    geometry(front); page_numbering(front, 1, 'lowerRoman'); footer_for(front)
    doc.add_paragraph('Document control', 'Title')
    table(doc, [['Item', 'Content'], ['Version', metadata['version']], ['Status', metadata['status']], ['Date', metadata['date']]] + controls)
    doc.add_paragraph('Purpose', 'Front Heading')
    rich(doc.add_paragraph(), metadata['purpose'])
    doc.add_page_break()
    doc.add_paragraph('Contents', 'Title')
    headings = [(level, title) for kind, *rest in blocks if kind == 'heading' for level, title in [rest]]
    native_toc(doc, headings)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    geometry(body); page_numbering(body, 1); footer_for(body)
    # Carry immediately preceding headings onto a dedicated figure page, avoiding
    # a nearly empty portrait page containing only a heading before the figure.
    plate_starts = {}
    for index, block in enumerate(blocks):
        if block[0] == 'image' and block[5]:
            start = index
            while start > 0 and blocks[start - 1][0] == 'heading': start -= 1
            plate_starts[start] = (index, block[5])
    chapter, plate_active, plate_end, figures = 0, False, -1, []
    for index, block in enumerate(blocks):
        kind = block[0]
        section_started = False
        if plate_active and index > plate_end and index not in plate_starts:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            geometry(sec)
            sec.different_first_page_header_footer = False
            page_numbering(sec)
            plate_active = False
            section_started = True
        if index in plate_starts:
            plate_end, paper = plate_starts[index]
            sec = doc.sections[-1] if index == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
            geometry(sec, paper)
            sec.different_first_page_header_footer = False
            if index != 0: page_numbering(sec)
            plate_active = True
        if kind == 'heading':
            chapter += 1
            p = doc.add_heading('', block[1])
            if block[1] == 1 and chapter > 1 and chapter_break and index not in plate_starts and not section_started:
                p.paragraph_format.page_break_before = True
            rich(p, block[2])
            p._p.append(element('bookmarkStart', id=chapter, name=f'chapter_{chapter}'))
            p._p.append(element('bookmarkEnd', id=chapter))
        elif kind == 'paragraph':
            p = doc.add_paragraph()
            if re.match(r'^Table\s*[A-Za-z]?\d+(?:[.．-]\d+)*\s+', block[1]):
                p.paragraph_format.keep_with_next = True
            rich(p, block[1])
        elif kind == 'list':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.55)
            p.paragraph_format.first_line_indent = Cm(-.45)
            rich(p, ('• ' if block[1] == '-' else block[1] + ' ') + block[2])
        elif kind == 'table':
            table(doc, block[1])
        elif kind == 'break':
            doc.add_page_break()
        elif kind == 'image':
            _, caption, path, width, height, paper = block
            max_width, max_height = {'a3': (37, 22), 'a4-landscape': (24.7, 13.5)}.get(paper, (16, 19))
            if paper:
                # Reserve real height for headings carried onto the same plate.
                start = next(start for start, (end, _) in plate_starts.items() if end == index)
                max_height -= sum({1: 1.8, 2: 1.5, 3: 1.3}[b[1]] for b in blocks[start:index])
                if max_height < 3: raise InputError('Too many headings before a figure page')
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = p.add_run().add_picture(str(path), width=Cm(min(max_width, max_height * width / height)))
            picture._inline.docPr.set('descr', caption)
            p = doc.add_paragraph(caption, 'Caption')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            figures.append({'caption': caption, 'source': str(path.relative_to(source.parent)) if path.is_relative_to(source.parent) else str(path), 'paper': paper or 'a4-inline'})
    doc.core_properties.title = metadata['title']
    doc.core_properties.subject = metadata['purpose'][:255]
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.exists() and not args.overwrite:
        raise InputError('Output already exists; use a new path or --overwrite explicitly')
    doc.save(args.output)
    receipt = {'status': 'BUILT_RENDER_AND_VISUAL_QA_REQUIRED', 'docx': args.output.name, 'sha256': hashlib.sha256(args.output.read_bytes()).hexdigest(), 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'attachment_links': doc.part._attachment_links, 'heading_count': len(headings), 'table_count': len(doc.tables), 'figures': figures, 'sections': len(doc.sections), 'body_font': args.body_font, 'heading_font': args.heading_font, 'font_availability': 'NOT_VALIDATED', 'toc': 'native_toc_h1_h2_requires_field_update', 'layout': {'chapter_page_break': chapter_break, 'cover_numbered': False, 'front_matter_numbering': 'lowerRoman', 'body_start': 1, 'footer_pt': 9}, 'toc_heading_count': sum(level <= 2 for level, _ in headings)}
    args.output.with_suffix('.build.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(receipt, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--markdown', required=True, type=Path)
    parser.add_argument('--metadata', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--body-font', default='Arial')
    parser.add_argument('--heading-font', default='Arial')
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    try:
        build(args)
    except (InputError, OSError, ValueError) as exc:
        parser.exit(2, f'build_report: {exc}\n')


if __name__ == '__main__':
    main()

# Attribution: adapted package copy of design-ai-project-architecture/scripts/build_report.py, Jiutian Lecture Lab, 2026-09-23.
